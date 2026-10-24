"""Turns the messy mentor sheet into PublicMentor records.

Rules: never guess a missing value, and never read Email, Contact or LinkedIn.
"""

import re
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.schemas import PublicMentor, Tools

NULL_TOKENS = {"", "na", "n/a", "-", "null", "none", "nan", "not available"}
TRUE_TOKENS = {"yes", "y", "true", "1"}
FALSE_TOKENS = {"no", "n", "false", "0"}

# Splits on ; , newline and " / ", but not inside brackets (the role "Business Generalist (...)").
_SPLIT = re.compile(r"(?:[;,\n]|\s/\s)(?![^()]*\))")
_SPLIT_BARE_SLASH = re.compile(r"(?:[;,\n/])(?![^()]*\))")
_BRACKETS = re.compile(r"\s*\([^)]*\)")
_COMMA = re.compile(r",(?![^()]*\))")


def _key(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.casefold())


def _canon(table: dict[str, list[str]]) -> dict[str, str]:
    return {_key(alias): name for name, aliases in table.items() for alias in aliases}


CITIES = _canon(
    {
        "Bengaluru": ["Bengaluru", "Bangalore", "Bangalore KA", "BLR"],
        "Mumbai": ["Mumbai", "Bombay"],
        "Delhi": ["Delhi", "New Delhi"],
        "Gurugram": ["Gurugram", "Gurgaon"],
    }
)
DOMAINS = _canon(
    {
        "Business Generalist": ["Business Generalist"],
        "Marketing": ["Marketing"],
        "Product Management": ["Product Management"],
        "B2B Sales and Partnerships/CSM": ["B2B sales and Partnerships/CSM"],
        "Growth": ["Growth"],
        "Venture Capital": ["Venture Capital"],
        "AI": ["AI"],
        "Data": ["Data"],
        "Category Management": ["Category Management"],
        "Behavioural": ["Behavioural"],
    }
)
SECTORS = _canon(
    {
        "Quick Commerce": ["Quick Commerce", "Quick comm"],
        "E-Commerce": ["E-Commerce", "Ecomm", "Ecom"],
        "Edtech": ["Edtech", "Edtech/Education", "Education"],
        "Healthtech": ["Healthtech", "Healthcare/ health tech", "Healthcare"],
        "Healthtech / Medtech": ["Healthtech / Medtech"],
        "HRTech": ["HRTech"],
        "SaaS": ["SaaS"],
        "Deeptech/AI": ["Deeptech/AI"],
        "Climatetech": ["Climatetech", "Climatech", "Climatetech/Greentech"],
        "Delivery/Logistics": ["Delivery/Logistics", "Delivery/Logisitics"],
        "Sportstech": ["Sportstech", "Sportstechg"],
        "Travel Tech": ["Travel tech"],
        "Tech, Info and Internet": ["Tech,Info and Internet"],
    }
)
ROLES = _canon(
    {
        "Founder's Office": ["Founders office", "Founder's Office"],
        "Performance Marketing": ["Perfomance Marketing", "Performance Marketing"],
        "Digital Marketing": ["Digital markeitng", "Digital Marketing"],
    }
)
AI_TOOLS = _canon(
    {
        "n8n": ["n8n"],
        "Zapier": ["Zapier"],
        "Make": ["Make"],
        "Notion AI": ["Notion AI"],
        "Factors.ai": ["Factors.ai"],
        "Google Ads": ["Google Ads"],
        "LinkedIn Ads": ["LinkedIn ads"],
    }
)

TOOL_COLUMNS = {
    "tableau": "Tableau",
    "hubspot": "Hubspot",
    "google_analytics": "Google Analytics",
    "jira": "Jira",
    "excel": "Excel",
    "mixpanel": "Mixpanel",
}


@dataclass
class CleaningReport:
    rows_read: int = 0
    mentors: int = 0
    fixes: Counter[str] = field(default_factory=Counter)

    def format(self) -> str:
        lines = [
            f"Rows read: {self.rows_read}",
            f"Unique mentors: {self.mentors}",
            *(f"  {name}: {count}" for name, count in sorted(self.fixes.items())),
        ]
        return "\n".join(lines)


def _text(value: Any, report: CleaningReport) -> str | None:
    """Stripped text, or None for blanks and placeholders like NA, '-' or 'Not available'."""
    if value is None:
        return None
    text = str(value).strip()
    if text.casefold() in NULL_TOKENS:
        if text:
            report.fixes["placeholders_set_to_null"] += 1
        return None
    return text


def _tidy(text: str) -> str:
    """Collapses spacing and fixes ALL CAPS or all lower case. Mixed case is left alone."""
    text = " ".join(text.split())
    if text.islower() or (text.isupper() and len(text) > 3):
        return text.title()
    return text


def _clean_one(text: str, table: dict[str, str]) -> str:
    return table.get(_key(text)) or _tidy(text)


def _clean_city(value: Any, report: CleaningReport) -> str | None:
    raw = _text(value, report)
    if raw is None:
        return None
    city = _clean_one(raw, CITIES)
    if city != raw:
        report.fixes["cities_normalised"] += 1
    return city


def _clean_single(value: Any, table: dict[str, str], report: CleaningReport) -> str | None:
    raw = _text(value, report)
    if raw is None:
        return None
    cleaned = _clean_one(raw, table)
    if cleaned != raw:
        report.fixes["text_fields_fixed"] += 1
    return cleaned


def _clean_domain(value: Any, report: CleaningReport) -> str:
    raw = _text(value, report)
    if raw is None:
        return "Unspecified"
    parts = [p.strip() for p in _COMMA.split(raw) if p.strip()]
    domain = ", ".join(_clean_one(p, DOMAINS) for p in parts)
    if domain != raw:
        report.fixes["domains_fixed"] += 1
    return domain


def _clean_list(
    value: Any, table: dict[str, str], report: CleaningReport, bare_slash: bool = False
) -> list[str]:
    raw = _text(value, report)
    if raw is None:
        return []
    pieces = (_SPLIT_BARE_SLASH if bare_slash else _SPLIT).split(raw)
    items: list[str] = []
    seen: set[str] = set()
    for piece in pieces:
        piece = piece.strip()
        if piece.casefold() in NULL_TOKENS:
            continue
        piece = _BRACKETS.sub(
            "", piece
        ).strip()  # drops example lists like "(Founder's Office/...)"
        name = _clean_one(piece, table)
        if name.casefold() not in seen:
            seen.add(name.casefold())
            items.append(name)
    if len(pieces) > 1:
        report.fixes["multi_value_cells_split"] += 1
    return items


def _clean_ai_tools(value: Any, report: CleaningReport) -> list[str]:
    tools = _clean_list(value, AI_TOOLS, report, bare_slash=True)
    names = [t for t in tools if t.casefold() not in TRUE_TOKENS | FALSE_TOKENS]
    # A bare "Yes" in this column is not a tool name.
    report.fixes["yes_no_text_dropped_from_ai_tools"] += len(tools) - len(names)
    return names


def _clean_bool(value: Any, report: CleaningReport) -> bool | None:
    raw = _text(value, report)
    if raw is None:
        return None
    token = raw.casefold()
    if token in TRUE_TOKENS or token in FALSE_TOKENS:
        if raw not in ("Yes", "No"):
            report.fixes["booleans_normalised"] += 1
        return token in TRUE_TOKENS
    report.fixes["unreadable_booleans_set_to_null"] += 1
    return None


def _clean_seniority(value: Any, report: CleaningReport) -> str | None:
    raw = _text(value, report)
    if raw is None:
        return None
    key = _key(raw)
    if key in ("midcareer", "midlevel"):
        result = "Mid-career"
    elif key == "senior":
        result = "Senior"
    else:
        report.fixes["unknown_seniority_set_to_null"] += 1
        return None
    if result != raw:
        report.fixes["seniority_normalised"] += 1
    return result


def _clean_quality(value: Any, report: CleaningReport) -> str | None:
    raw = _text(value, report)
    if raw is None:
        return None
    result = {"excellent": "Excellent", "good": "Good"}.get(raw.casefold())
    if result is None:
        report.fixes["unknown_quality_set_to_null"] += 1
    elif result != raw:
        report.fixes["quality_normalised"] += 1
    return result


def _clean_activities(value: Any, report: CleaningReport) -> int | None:
    raw = _text(value, report)
    if raw is None:
        return None
    match = re.fullmatch(r"(\d+)(?:\.0+)?(?:\s*activit(?:y|ies))?", raw, re.IGNORECASE)
    if match is None:
        report.fixes["unreadable_activities_set_to_null"] += 1
        return None
    if not isinstance(value, (int, float)):
        report.fixes["activities_parsed"] += 1
    return int(match.group(1))


def _to_mentor(row: dict[str, Any], report: CleaningReport) -> PublicMentor:
    name = _clean_single(row.get("Name"), {}, report)
    return PublicMentor(
        id=str(row["Mentor ID"]).strip().upper(),
        name=name or str(row["Mentor ID"]).strip(),
        city=_clean_city(row.get("City"), report),
        domain=_clean_domain(row.get("Domain"), report),
        company=_clean_single(row.get("Company"), {}, report),
        sector=_clean_single(row.get("Master Sector"), SECTORS, report),
        sectors=_clean_list(row.get("Sectors of Expertise"), SECTORS, report),
        roles=_clean_list(row.get("Roles of Expertise"), ROLES, report),
        seniority=_clean_seniority(row.get("Seniority"), report),
        built_ai=_clean_bool(row.get("Built AI Product/Automation/Workflow"), report),
        ai_tools=_clean_ai_tools(row.get("AI Tools of Expertise"), report),
        tools=Tools(**{k: _clean_bool(row.get(col), report) for k, col in TOOL_COLUMNS.items()}),
        quality=_clean_quality(row.get("Quality of Mentorship"), report),
        activities=_clean_activities(row.get("Activities Participated"), report),
    )


def _filled(row: dict[str, Any]) -> int:
    return sum(1 for v in row.values() if v is not None and str(v).strip())


def _dedupe(rows: Iterable[dict[str, Any]], report: CleaningReport) -> list[dict[str, Any]]:
    """One row per Mentor ID. Keeps the fullest row, the first one on a tie."""
    best: dict[str, dict[str, Any]] = {}
    for row in rows:
        mentor_id = str(row["Mentor ID"]).strip().upper()
        if mentor_id not in best:
            best[mentor_id] = row
            continue
        report.fixes["duplicate_rows_removed"] += 1
        if _filled(row) > _filled(best[mentor_id]):
            best[mentor_id] = row
    return sorted(best.values(), key=lambda r: str(r["Mentor ID"]).strip().upper())


def read_rows(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Mentors"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    records = [dict(zip(header, values, strict=False)) for values in rows]
    workbook.close()
    return [r for r in records if r.get("Mentor ID")]


def clean_mentors(path: Path) -> tuple[list[PublicMentor], CleaningReport]:
    report = CleaningReport()
    rows = read_rows(path)
    report.rows_read = len(rows)
    mentors = [_to_mentor(row, report) for row in _dedupe(rows, report)]
    report.mentors = len(mentors)
    return mentors, report
